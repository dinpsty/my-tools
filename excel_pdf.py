import io
import os
import posixpath
import shutil
import subprocess
import tempfile
import zipfile
from pathlib import Path
import json
from openpyxl import load_workbook
from lxml import etree


NS_MAIN = (
    "http://schemas.openxmlformats.org/"
    "spreadsheetml/2006/main"
)
NS_REL = (
    "http://schemas.openxmlformats.org/"
    "officeDocument/2006/relationships"
)
NS_PKG = (
    "http://schemas.openxmlformats.org/"
    "package/2006/relationships"
)

MAX_UPLOAD = 10 * 1024 * 1024
MAX_UNCOMPRESSED = 80 * 1024 * 1024
MAX_ENTRIES = 1500
MAX_PDF = 30 * 1024 * 1024


def validasi_excel(data):
    if not data:
        raise ValueError("File Excel kosong.")

    if len(data) > MAX_UPLOAD:
        raise ValueError("Ukuran Excel maksimal 10 MB.")

    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            daftar = z.infolist()

            if len(daftar) > MAX_ENTRIES:
                raise ValueError("Komponen Excel terlalu banyak.")

            ukuran = sum(x.file_size for x in daftar)

            if ukuran > MAX_UNCOMPRESSED:
                raise ValueError("Isi Excel terlalu besar.")

            if z.testzip() is not None:
                raise ValueError("File Excel rusak.")

            nama = set(z.namelist())

            if "xl/workbook.xml" not in nama:
                raise ValueError("File bukan XLSX valid.")

            if "xl/_rels/workbook.xml.rels" not in nama:
                raise ValueError("Relasi workbook tidak ada.")

            if any(
                x.startswith("xl/externalLinks/")
                for x in nama
            ):
                raise ValueError(
                    "Workbook memiliki external links. "
                    "Putuskan tautan eksternal sebelum upload."
                )

    except zipfile.BadZipFile:
        raise ValueError("File XLSX tidak valid.")


def baca_xml(data):
    parser = etree.XMLParser(
        resolve_entities=False,
        no_network=True,
        huge_tree=False
    )
    return etree.fromstring(data, parser=parser)


def pilih_sheet_excel(data, nama_sheet):
    validasi_excel(data)

    sumber_buffer = io.BytesIO(data)
    hasil_buffer = io.BytesIO()

    with zipfile.ZipFile(sumber_buffer) as sumber:
        workbook = baca_xml(
            sumber.read("xl/workbook.xml")
        )

        relasi = baca_xml(
            sumber.read("xl/_rels/workbook.xml.rels")
        )

        sheet_tersedia = []

        for elemen in workbook.findall(
            f".//{{{NS_MAIN}}}sheets/{{{NS_MAIN}}}sheet"
        ):
            sheet_tersedia.append(elemen)

        ditemukan = False
        index_aktif = 0

        for index, sheet in enumerate(sheet_tersedia):
            if sheet.get("name") == nama_sheet:
                sheet.set("state", "visible")
                ditemukan = True
                index_aktif = index
            else:
                sheet.set("state", "hidden")

        if not ditemukan:
            raise ValueError("Sheet tidak ditemukan.")

        views = workbook.find(f"{{{NS_MAIN}}}bookViews")

        if views is not None:
            for view in views:
                if view.tag == f"{{{NS_MAIN}}}workbookView":
                    view.set("activeTab", str(index_aktif))
                    view.set("firstSheet", str(index_aktif))

        rel_map = {}

        for rel in relasi.findall(
            f"{{{NS_PKG}}}Relationship"
        ):
            target = rel.get("Target", "")

            if target.startswith("/"):
                path = target.lstrip("/")
            else:
                path = posixpath.normpath(
                    posixpath.join("xl", target)
                )

            rel_map[rel.get("Id")] = path

        sheet_target = sheet_tersedia[index_aktif]
        rid = sheet_target.get(f"{{{NS_REL}}}id")
        path_sheet = rel_map.get(rid, "")

        if (
            not path_sheet.startswith("xl/worksheets/")
            or path_sheet not in sumber.namelist()
        ):
            raise ValueError(
                "Sheet Excel tidak dapat diproses."
            )

        with zipfile.ZipFile(
            hasil_buffer,
            "w",
            zipfile.ZIP_DEFLATED
        ) as hasil:
            for info in sumber.infolist():
                isi = sumber.read(info.filename)

                if info.filename == "xl/workbook.xml":
                    isi = etree.tostring(
                        workbook,
                        encoding="UTF-8",
                        xml_declaration=True
                    )

                hasil.writestr(info, isi)

    return hasil_buffer.getvalue()


def konversi_pdf(data, nomor_sheet):
    executable = (
        shutil.which("libreoffice")
        or shutil.which("soffice")
    )

    if not executable:
        raise RuntimeError(
            "LibreOffice tidak tersedia di server."
        )

    pengaturan = json.dumps({
        "SheetRange": {
            "type": "string",
            "value": str(nomor_sheet)
        },
        "UseLosslessCompression": {
            "type": "boolean",
            "value": "true"
        },
        "ReduceImageResolution": {
            "type": "boolean",
            "value": "false"
        }
    })

    with tempfile.TemporaryDirectory() as tmp:
        folder = Path(tmp)
        input_file = folder / "input.xlsx"
        output_folder = folder / "output"
        profile = folder / "profile"

        output_folder.mkdir()
        input_file.write_bytes(data)

        command = [
            executable,
            f"-env:UserInstallation={profile.as_uri()}",
            "--headless",
            "--convert-to",
            f"pdf:calc_pdf_Export:{pengaturan}",
            "--outdir",
            str(output_folder),
            str(input_file)
        ]

        env = os.environ.copy()
        env["HOME"] = str(folder)
        env["TMPDIR"] = str(folder)

        try:
            proses = subprocess.run(
                command,
                capture_output=True,
                text=True,
                timeout=90,
                cwd=str(folder),
                env=env
            )
        except subprocess.TimeoutExpired:
            raise RuntimeError(
                "Konversi terlalu lama dan dihentikan."
            )

        pdf_path = output_folder / "input.pdf"

        if proses.returncode != 0 or not pdf_path.exists():
            raise RuntimeError(
                "Gagal membuat PDF: "
                + (proses.stderr or proses.stdout)[:500]
            )

        if pdf_path.stat().st_size == 0:
            raise RuntimeError("File PDF kosong.")

        if pdf_path.stat().st_size > MAX_PDF:
            raise ValueError("Hasil PDF terlalu besar.")

        return pdf_path.read_bytes()


def konversi_satu_sheet(data, nama_sheet):
    validasi_excel(data)

    workbook = load_workbook(
        io.BytesIO(data),
        read_only=True,
        keep_links=False
    )

    try:
        if nama_sheet not in workbook.sheetnames:
            raise ValueError("Sheet tidak ditemukan.")

        nomor_sheet = (
            workbook.sheetnames.index(nama_sheet) + 1
        )

    finally:
        workbook.close()

    return konversi_pdf(data, nomor_sheet)
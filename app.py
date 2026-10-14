import io

import os

import re

import shutil

import subprocess

import tempfile

import zipfile

import posixpath

from pathlib import Path

import streamlit as st

from docx import Document

from openpyxl import Workbook, load_workbook

from lxml import etree

from pypdf import PdfReader, PdfWriter

from reportlab.pdfgen import canvas

st.set_page_config(

    page_title="My Tools",

    page_icon="🛠️",

    layout="centered"

)

st.title("🛠️ MY TOOLS")

pilihan = st.selectbox(

    "Pilih Tools",

    [

        "PDF Blank",

        "DOC Blank",

        "XCL Blank",

        "XCL Sheet to PDF",

        "DOC Page to PDF",

        "XCL Sheet to PDF (All Sheets)",

        "DOC Page to PDF (All Pages)"

    ]

)

karakter_terlarang = r'[<>:"/\\|?*]'

def bersihkan_nama(nama):

    nama = re.sub(karakter_terlarang, "", nama)

    return nama.strip().rstrip(". ")

def daftar_nama_bersih(teks, ekstensi):

    hasil = []

    sudah_ada = set()

    for nama in teks.splitlines():

        nama = nama.strip()

        if nama.lower().endswith(ekstensi):

            nama = nama[:-len(ekstensi)]

        nama = bersihkan_nama(nama)

        if nama.lower().endswith(" (blank)"):

            nama = nama[:-8].strip()

        if nama and nama.lower() not in sudah_ada:

            sudah_ada.add(nama.lower())

            hasil.append(nama)

    return hasil

def buat_zip(daftar_file):

    buffer = io.BytesIO()

    with zipfile.ZipFile(

        buffer, "w", zipfile.ZIP_DEFLATED

    ) as zip_file:

        for nama, data in daftar_file:

            zip_file.writestr(nama, data)

    return buffer.getvalue()

def nama_unik(nama, digunakan):

    nama_awal = nama

    nomor = 2

    while nama.lower() in digunakan:

        nama = f"{nama_awal} ({nomor})"

        nomor += 1

    digunakan.add(nama.lower())

    return nama

def cari_libreoffice():

    lokasi = (

        shutil.which("libreoffice")

        or shutil.which("soffice")

    )

    if lokasi:

        return lokasi

    lokasi_windows = (

        r"C:\Program Files\LibreOffice\program\soffice.exe"

    )

    if os.path.exists(lokasi_windows):

        return lokasi_windows

    raise RuntimeError(

        "LibreOffice belum tersedia di server."

    )

def konversi_ke_pdf(data, ekstensi):

    with tempfile.TemporaryDirectory() as folder:

        path_input = Path(folder) / f"input{ekstensi}"

        path_output = Path(folder) / "hasil"

        path_profile = Path(folder) / "profile"

        path_output.mkdir()

        path_input.write_bytes(data)

        perintah = [

            cari_libreoffice(),

            f"-env:UserInstallation={path_profile.as_uri()}",

            "--headless",

            "--convert-to",

            "pdf",

            "--outdir",

            str(path_output),

            str(path_input)

        ]

        proses = subprocess.run(

            perintah,

            capture_output=True,

            text=True,

            timeout=180

        )

        path_pdf = path_output / "input.pdf"

        if proses.returncode != 0 or not path_pdf.exists():

            raise RuntimeError(

                proses.stderr

                or proses.stdout

                or "Gagal mengonversi file."

            )

        return path_pdf.read_bytes()

def siapkan_excel_sheet(data_excel, nama_sheet, mode, jumlah_halaman=0):
    ns_main = (
        "http://schemas.openxmlformats.org/"
        "spreadsheetml/2006/main"
    )
    ns_rel = (
        "http://schemas.openxmlformats.org/"
        "officeDocument/2006/relationships"
    )
    ns_pkg = (
        "http://schemas.openxmlformats.org/"
        "package/2006/relationships"
    )

    input_buffer = io.BytesIO(data_excel)
    output_buffer = io.BytesIO()

    with zipfile.ZipFile(input_buffer, "r") as sumber:
        workbook_xml = etree.fromstring(
            sumber.read("xl/workbook.xml")
        )

        relationships_xml = etree.fromstring(
            sumber.read("xl/_rels/workbook.xml.rels")
        )

        hubungan = {}

        for rel in relationships_xml.findall(
            f"{{{ns_pkg}}}Relationship"
        ):
            target = rel.get("Target", "")

            if target.startswith("/"):
                path = target.lstrip("/")
            else:
                path = posixpath.normpath(
                    posixpath.join("xl", target)
                )

            hubungan[rel.get("Id")] = path

        target_sheet = None
        indeks_sheet = None

        sheets = workbook_xml.findall(
            f".//{{{ns_main}}}sheets/{{{ns_main}}}sheet"
        )

        for indeks, sheet in enumerate(sheets):
            if sheet.get("name") == nama_sheet:
                sheet.set("state", "visible")
                indeks_sheet = indeks

                rel_id = sheet.get(f"{{{ns_rel}}}id")
                target_sheet = hubungan.get(rel_id)

            else:
                sheet.set("state", "hidden")

        if target_sheet is None:
            raise ValueError(
                f"Sheet '{nama_sheet}' tidak ditemukan."
            )

        views = workbook_xml.find(
            f"{{{ns_main}}}bookViews"
        )

        if views is not None and indeks_sheet is not None:
            for view in views:
                if view.tag == f"{{{ns_main}}}workbookView":
                    view.set("activeTab", str(indeks_sheet))

        with zipfile.ZipFile(
            output_buffer, "w"
        ) as hasil:

            for info in sumber.infolist():
                data = sumber.read(info.filename)

                if info.filename == "xl/workbook.xml":
                    data = etree.tostring(
                        workbook_xml,
                        encoding="UTF-8",
                        xml_declaration=True
                    )

                elif (
                    info.filename == target_sheet
                    and mode in [
                        "Auto Layout",
                        "Fit to Width",
                        "Fit to Pages",
                    ]
                ):
                    root = etree.fromstring(data)

                    sheet_pr = root.find(
                        f"{{{ns_main}}}sheetPr"
                    )

                    if sheet_pr is None:
                        sheet_pr = etree.Element(
                            f"{{{ns_main}}}sheetPr"
                        )
                        root.insert(0, sheet_pr)

                    setup_pr = sheet_pr.find(
                        f"{{{ns_main}}}pageSetUpPr"
                    )

                    if setup_pr is None:
                        setup_pr = etree.SubElement(
                            sheet_pr,
                            f"{{{ns_main}}}pageSetUpPr"
                        )

                    setup_pr.set("fitToPage", "1")

                    page_setup = root.find(
                        f"{{{ns_main}}}pageSetup"
                    )

                    if page_setup is None:
                        page_setup = etree.Element(
                            f"{{{ns_main}}}pageSetup"
                        )

                        posisi = len(root)

                        for i, element in enumerate(root):
                            tag = etree.QName(
                                element
                            ).localname

                            if tag in [
                                "headerFooter",
                                "rowBreaks",
                                "colBreaks",
                                "customProperties",
                                "cellWatches",
                                "ignoredErrors",
                                "smartTags",
                                "drawing",
                                "legacyDrawing",
                                "legacyDrawingHF",
                                "picture",
                                "oleObjects",
                                "controls",
                                "webPublishItems",
                                "tableParts",
                                "extLst"
                            ]:
                                posisi = i
                                break

                        root.insert(posisi, page_setup)

                    page_setup.set("fitToWidth", "1")

                    if mode == "Fit to Pages":
                        page_setup.set("fitToHeight", str(jumlah_halaman))
                    else:
                        page_setup.set("fitToHeight", "0")

                    if "scale" in page_setup.attrib:
                        del page_setup.attrib["scale"]

                    data = etree.tostring(
                        root,
                        encoding="UTF-8",
                        xml_declaration=True
                    )

                hasil.writestr(info, data)

    return output_buffer.getvalue()


def konversi_sheet(data_excel, nama_sheet, mode, jumlah_halaman=0):
    data_sheet = siapkan_excel_sheet(
        data_excel,
        nama_sheet,
        mode,
        jumlah_halaman
    )

    return konversi_ke_pdf(
        data_sheet,
        ".xlsx"
    )

def pdf_satu_halaman(pdf_bytes, nomor_halaman):

    reader = PdfReader(io.BytesIO(pdf_bytes))

    writer = PdfWriter()

    writer.add_page(reader.pages[nomor_halaman - 1])

    buffer = io.BytesIO()

    writer.write(buffer)

    return buffer.getvalue()

@st.cache_data(show_spinner=False)

def konversi_word_cached(data):

    return konversi_ke_pdf(data, ".docx")

def proses_excel(file_excel, semua_sheet=False):
    data_excel = file_excel.getvalue()

    try:
        workbook = load_workbook(
            io.BytesIO(data_excel),
            read_only=True
        )

        daftar_sheet = workbook.sheetnames
        workbook.close()

    except Exception as e:
        st.error(f"Gagal membaca Excel: {e}")
        return

    st.write(f"Jumlah sheet: {len(daftar_sheet)}")

    if semua_sheet:
        sheet_terpilih = daftar_sheet
        st.write("Semua sheet akan dikonversi.")
        st.write(", ".join(daftar_sheet))

    else:
        sheet_terpilih = st.multiselect(
            "Pilih Sheet",
            daftar_sheet
        )

    mode = st.radio(
        "Metode Konversi",
        [
            "Auto Layout",
            "Original Layout",
            "Fit to Width",
            "Fit to Pages"
        ],
        index=0,
        key=f"mode_excel_{semua_sheet}"
    )

    jumlah_halaman = 0

    if mode == "Fit to Pages":
        jumlah_halaman = st.number_input(
            "Maksimal halaman vertikal",
            min_value=1,
            max_value=100,
            value=2,
            step=1,
            key=f"halaman_excel_{semua_sheet}"
        )

    if st.button("Convert ke PDF"):
        if not sheet_terpilih:
            st.warning("Pilih minimal satu sheet.")
            return

        hasil_file = []
        gagal = []
        digunakan = set()

        progress = st.progress(0)

        for index, nama_sheet in enumerate(sheet_terpilih):
            try:
                pdf_bytes = konversi_sheet(
                    data_excel,
                    nama_sheet,
                    mode,
                    jumlah_halaman
                )

                nama = bersihkan_nama(nama_sheet)
                nama = nama_unik(
                    nama or "Sheet",
                    digunakan
                )

                hasil_file.append(
                    (f"{nama}.pdf", pdf_bytes)
                )

            except Exception as e:
                gagal.append(
                    f"{nama_sheet}: {e}"
                )

            progress.progress(
                (index + 1) / len(sheet_terpilih)
            )

        if hasil_file:
            st.success(
                f"{len(hasil_file)} sheet berhasil dikonversi."
            )

            if len(hasil_file) == 1:
                nama, data = hasil_file[0]

                st.download_button(
                    "⬇️ Download PDF",
                    data=data,
                    file_name=nama,
                    mime="application/pdf"
                )

            else:
                st.download_button(
                    "⬇️ Download ZIP",
                    data=buat_zip(hasil_file),
                    file_name="excel_sheet_pdf.zip",
                    mime="application/zip"
                )

        if gagal:
            st.error(
                "Beberapa sheet gagal dikonversi."
            )

            for pesan in gagal:
                st.write(pesan)

def proses_word(file_word, semua_halaman=False):

    data_word = file_word.getvalue()

    try:

        with st.spinner("Membaca halaman Word..."):

            pdf_bytes = konversi_word_cached(data_word)

        reader = PdfReader(io.BytesIO(pdf_bytes))

        jumlah_halaman = len(reader.pages)

    except Exception as e:

        st.error(f"Gagal membaca dokumen: {e}")

        return

    st.write(f"Jumlah halaman: {jumlah_halaman}")

    if semua_halaman:

        halaman_terpilih = list(

            range(1, jumlah_halaman + 1)

        )

        st.write("Semua halaman akan dikonversi.")

    else:

        halaman_terpilih = st.multiselect(

            "Pilih Halaman",

            list(range(1, jumlah_halaman + 1)),

            format_func=lambda x: f"Halaman {x}"

        )

    if st.button("Convert ke PDF"):

        if not halaman_terpilih:

            st.warning("Pilih minimal satu halaman.")

            return

        nama_dokumen = bersihkan_nama(

            Path(file_word.name).stem

        ) or "Dokumen"

        hasil_file = []

        for nomor in sorted(halaman_terpilih):

            data_halaman = pdf_satu_halaman(

                pdf_bytes, nomor

            )

            nama_pdf = (

                f"{nama_dokumen} - Page {nomor:03d}.pdf"

            )

            hasil_file.append(

                (nama_pdf, data_halaman)

            )

        st.success(

            f"{len(hasil_file)} halaman berhasil dibuat."

        )

        if len(hasil_file) == 1:

            nama, data = hasil_file[0]

            st.download_button(

                "⬇️ Download PDF",

                data=data,

                file_name=nama,

                mime="application/pdf"

            )

        else:

            st.download_button(

                "⬇️ Download ZIP",

                data=buat_zip(hasil_file),

                file_name="word_pages_pdf.zip",

                mime="application/zip"

            )

if pilihan in [

    "PDF Blank",

    "DOC Blank",

    "XCL Blank"

]:

    st.subheader(pilihan)

    st.write("Masukkan nama file satu per baris.")

    input_nama = st.text_area(

        "Nama File",

        height=250,

        placeholder="Contoh:\nAndi\nBudi\nCitra"

    )

    if pilihan == "PDF Blank":

        ekstensi = ".pdf"

    elif pilihan == "DOC Blank":

        ekstensi = ".docx"

    else:

        ekstensi = ".xlsx"

    if st.button("Buat File"):

        daftar_nama = daftar_nama_bersih(

            input_nama, ekstensi

        )

        if not daftar_nama:

            st.error("Masukkan minimal satu nama file.")

        else:

            daftar_file = []

            for nama in daftar_nama:

                buffer = io.BytesIO()

                if pilihan == "PDF Blank":

                    pdf = canvas.Canvas(buffer)

                    pdf.showPage()

                    pdf.save()

                elif pilihan == "DOC Blank":

                    document = Document()

                    document.save(buffer)

                else:

                    workbook = Workbook()

                    workbook.save(buffer)

                    workbook.close()

                nama_file = f"{nama} (BLANK){ekstensi}"

                daftar_file.append(

                    (nama_file, buffer.getvalue())

                )

            st.success(

                f"{len(daftar_file)} file berhasil dibuat."

            )

            st.download_button(

                "⬇️ Download File",

                data=buat_zip(daftar_file),

                file_name="blank_files.zip",

                mime="application/zip"

            )

elif pilihan == "XCL Sheet to PDF":

    st.subheader("📊 XCL Sheet to PDF")

    st.write(

        "Upload Excel dan pilih sheet yang ingin "

        "dijadikan PDF."

    )

    file_excel = st.file_uploader(

        "Upload Excel",

        type=["xlsx"]

    )

    if file_excel:

        proses_excel(file_excel)

elif pilihan == "XCL Sheet to PDF (All Sheets)":

    st.subheader("📊 XCL Sheet to PDF (All Sheets)")

    st.write(

        "Seluruh sheet Excel akan dikonversi "

        "menjadi PDF terpisah."

    )

    file_excel = st.file_uploader(

        "Upload Excel",

        type=["xlsx"]

    )

    if file_excel:

        proses_excel(file_excel, semua_sheet=True)

elif pilihan == "DOC Page to PDF":

    st.subheader("📝 DOC Page to PDF")

    st.write(

        "Upload Word dan pilih halaman yang "

        "ingin dijadikan PDF."

    )

    file_word = st.file_uploader(

        "Upload Word",

        type=["docx"]

    )

    if file_word:

        proses_word(file_word)

elif pilihan == "DOC Page to PDF (All Pages)":

    st.subheader("📝 DOC Page to PDF (All Pages)")

    st.write(

        "Seluruh halaman Word akan disimpan "

        "sebagai PDF terpisah."

    )

    file_word = st.file_uploader(

        "Upload Word",

        type=["docx"]

    )

    if file_word:

        proses_word(file_word, semua_halaman=True)